# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# ─── Page margins ──────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)

style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(13)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

# ─── Helper functions ──────────────────────────────────────────────
def add_run(paragraph, text, bold=False, italic=False, size=None, color=None, underline=False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if underline:
        run.underline = True
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    run.font.name = 'Times New Roman'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    return run

def add_heading_custom(text, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=3):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    add_run(p, text, bold=bold, size=size)
    return p

def add_body(text, bold=False, italic=False, indent=0, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    add_run(p, text, bold=bold, italic=italic)
    return p

def add_bullet(text, indent=1, bold_prefix="", space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    if bold_prefix:
        add_run(p, "•  ", bold=False)
        add_run(p, bold_prefix, bold=True)
        add_run(p, text)
    else:
        add_run(p, "•  " + text)
    return p

# ═══════════════════════════════════════════════════════════════════
# HEADER
# ═══════════════════════════════════════════════════════════════════
table_header = doc.add_table(rows=1, cols=2)
table_header.alignment = WD_TABLE_ALIGNMENT.CENTER
cell_left = table_header.cell(0, 0)
cell_right = table_header.cell(0, 1)

p_left = cell_left.paragraphs[0]
p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(p_left, "HỌC VIỆN CÔNG NGHỆ\nBƯU CHÍNH VIỄN THÔNG", bold=True, size=12)

p_right = cell_right.paragraphs[0]
p_right.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(p_right, "CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, size=12)
p_right2 = cell_right.add_paragraph()
p_right2.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(p_right2, "Độc lập - Tự do - Hạnh phúc", bold=True, italic=True, size=12, underline=True)

p_khoa = cell_left.add_paragraph()
p_khoa.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_run(p_khoa, "KHOA CÔNG NGHỆ THÔNG TIN 1", bold=True, size=11)

# Remove table borders
for row in table_header.rows:
    for cell in row.cells:
        tc = cell._element
        tcPr = tc.get_or_add_tcPr()
        tcBorders = tcPr.makeelement(qn('w:tcBorders'), {})
        for border_name in ['top', 'left', 'bottom', 'right']:
            border = tcBorders.makeelement(qn(f'w:{border_name}'), {
                qn('w:val'): 'none', qn('w:sz'): '0', qn('w:space'): '0', qn('w:color'): 'auto'
            })
            tcBorders.append(border)
        tcPr.append(tcBorders)

# Title
doc.add_paragraph()
add_heading_custom("BÁO CÁO 2", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12)
add_heading_custom("THỰC TẬP CƠ SỞ - NHÓM 7", size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=12)

# Student info
info_lines = [
    ("Họ và tên: ", "Lý Hải Quân", "Mã SV: ", "B23DCCN676"),
    ("Email: ", "lyhaiquan020705@gmail.com", "Số điện thoại: ", "0976594856"),
    ("Ngành đào tạo: ", "Công nghệ Thông tin", "Hệ đào tạo: ", "Đại học chính quy"),
]
for left_label, left_val, right_label, right_val in info_lines:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    add_run(p, left_label, bold=True, size=13)
    add_run(p, left_val, size=13)
    add_run(p, "                    ", size=13)
    add_run(p, right_label, bold=True, size=13)
    add_run(p, right_val, size=13)

doc.add_paragraph()

# ═══════════════════════════════════════════════════════════════════
# A/ CÁC NỘI DUNG THỰC TẬP
# ═══════════════════════════════════════════════════════════════════
add_heading_custom("A/ Các nội dung thực tập", size=14)

add_heading_custom("1. Tên dự án thực hiện", size=13)
add_body("Smart AI Chatbot – Trợ lý nhóm chat thông minh tích hợp Gemini AI.", indent=0.5)

add_heading_custom("2. Mục tiêu của dự án", size=13)
add_bullet("Xây dựng một ứng dụng web chatbot cho phép người dùng trò chuyện với trợ lý ảo AI (Gemini 2.5 Flash), hỗ trợ tự động dịch ngôn ngữ, phản hồi thông minh theo ngữ cảnh và tóm tắt cuộc trò chuyện.", bold_prefix="Mục tiêu chính: ")
add_bullet("", bold_prefix="Mục tiêu kỹ thuật:")
add_bullet("Tìm hiểu kiến trúc Full-Stack: Backend (Node.js + Express) và Frontend (React + Vite).", indent=1.8)
add_bullet("Tích hợp API của Google Generative AI (Gemini 2.5 Flash) vào ứng dụng thực tế.", indent=1.8)
add_bullet("Xử lý giới hạn tốc độ API (Rate Limiting) với cơ chế bắt lỗi 429.", indent=1.8)
add_bullet("Container hóa toàn bộ dự án bằng Docker và Docker Compose.", indent=1.8)
add_bullet("Xây dựng giao diện người dùng hiện đại, thân thiện và responsive.", indent=1.8)

# ═══════════════════════════════════════════════════════════════════
# B/ NỘI DUNG THỰC HIỆN
# ═══════════════════════════════════════════════════════════════════
doc.add_paragraph()
add_heading_custom("B/ Nội dung thực hiện nhiệm vụ được phân công", size=14)

# ─── 1. Giới thiệu công nghệ ──────────────────────────────────────
add_heading_custom("1. Giới thiệu và Lý do lựa chọn công nghệ", size=13)

add_body("Dự án sử dụng kiến trúc Client-Server hiện đại với các công nghệ sau:", italic=True, indent=0.5)

# Backend techs
add_heading_custom("a) Backend – Node.js + Express", size=13, space_before=4)
add_bullet("Nền tảng runtime JavaScript phía server, cho phép xử lý I/O bất đồng bộ hiệu quả, phù hợp với ứng dụng real-time.", bold_prefix="Node.js: ")
add_bullet("Framework web tối giản, giúp tạo RESTful API endpoint nhanh chóng với cú pháp rõ ràng.", bold_prefix="Express.js: ")
add_bullet("Middleware cho phép Frontend (chạy trên port khác) gọi API đến Backend một cách an toàn.", bold_prefix="CORS: ")
add_bullet("Quản lý biến môi trường (.env), giữ API Key và cấu hình nhạy cảm tách biệt khỏi source code.", bold_prefix="Dotenv: ")
add_bullet("SDK chính thức của Google để tương tác với Gemini API, hỗ trợ chat đa lượt (multi-turn conversation) với lịch sử hội thoại.", bold_prefix="@google/generative-ai: ")

# Frontend techs
add_heading_custom("b) Frontend – React + Vite", size=13, space_before=4)
add_bullet("Thư viện xây dựng giao diện theo component, sử dụng Hooks (useState, useRef, useEffect) để quản lý state và side effects.", bold_prefix="React: ")
add_bullet("Build tool thế hệ mới, khởi động dev server gần như tức thì nhờ ES modules, tốc độ build production nhanh gấp nhiều lần Webpack.", bold_prefix="Vite: ")
add_bullet("HTTP client hỗ trợ Promise, interceptor và xử lý lỗi tốt hơn so với fetch API gốc.", bold_prefix="Axios: ")

# AI
add_heading_custom("c) Mô hình AI – Gemini 2.5 Flash", size=13, space_before=4)
add_bullet("Mô hình ngôn ngữ lớn (LLM) mới nhất của Google, tối ưu cho tốc độ phản hồi nhanh với chi phí thấp.", bold_prefix="Gemini 2.5 Flash: ")
add_bullet("Cấu hình hành vi AI ngay từ lúc khởi tạo model, đảm bảo AI tuân thủ nghiêm ngặt 3 quy tắc: tự động dịch, phản hồi thông minh, tóm tắt hội thoại.", bold_prefix="System Instruction: ")
add_bullet("API hỗ trợ gửi kèm lịch sử hội thoại (history) để AI hiểu ngữ cảnh xuyên suốt cuộc trò chuyện.", bold_prefix="Multi-turn Chat: ")

# Docker
add_heading_custom("d) Container hóa – Docker + Docker Compose", size=13, space_before=4)
add_bullet("Đóng gói ứng dụng cùng toàn bộ dependencies vào container, đảm bảo chạy nhất quán trên mọi môi trường (dev, staging, production).", bold_prefix="Docker: ")
add_bullet("Frontend sử dụng 2 giai đoạn: build bằng Node Alpine → serve bằng Nginx Alpine, giảm kích thước image đáng kể.", bold_prefix="Multi-stage Build: ")
add_bullet("Điều phối 2 services (backend + frontend) chỉ với một lệnh duy nhất, kèm quản lý biến môi trường an toàn qua file .env.", bold_prefix="Docker Compose: ")
add_bullet("Frontend container sử dụng Nginx để vừa serve file tĩnh React, vừa proxy các request /api tới backend container thông qua Docker internal network.", bold_prefix="Nginx Reverse Proxy: ")

# ─── 2. Kiến trúc hệ thống ────────────────────────────────────────
doc.add_paragraph()
add_heading_custom("2. Kiến trúc hệ thống", size=13)

add_body("Hệ thống được thiết kế theo mô hình 3 tầng:", indent=0.5)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(6)
p.paragraph_format.left_indent = Cm(0.5)
add_run(p, """
┌─────────────────────────────────────────────────────────┐
│                    Docker Compose                       │
│                                                         │
│  ┌──────────────┐         ┌──────────────────────────┐  │
│  │   Frontend   │  /api   │        Backend           │  │
│  │  (Nginx:80)  │───────► │   (Express:5000)         │  │
│  │  React SPA   │         │   @google/generative-ai  │  │
│  └──────────────┘         └────────────┬─────────────┘  │
│                                        │                │
│                                        ▼                │
│                              ┌──────────────────┐       │
│                              │   Gemini 2.5     │       │
│                              │   Flash API      │       │
│                              └──────────────────┘       │
└─────────────────────────────────────────────────────────┘
""", size=9)

# ─── 3. Chi tiết triển khai ────────────────────────────────────────
add_heading_custom("3. Chi tiết triển khai", size=13, space_before=8)

add_heading_custom("3.1. Backend (Node.js + Express)", size=13, space_before=6)

add_body("Cấu trúc thư mục:", bold=True, indent=0.5, space_after=2)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Cm(1)
p.paragraph_format.space_after = Pt(4)
add_run(p, "backend/\n├── .env              (Biến môi trường: PORT, GEMINI_API_KEY)\n├── .dockerignore\n├── Dockerfile\n├── index.js          (Server chính)\n├── package.json\n└── node_modules/", size=11)

add_body("Các thành phần chính trong index.js:", bold=True, indent=0.5, space_after=2)
add_bullet("Khởi tạo Express server với middleware CORS và JSON parser, lắng nghe trên port 5000.")
add_bullet("Khởi tạo Google Generative AI SDK với systemInstruction chứa 3 quy tắc hành vi cho chatbot.")
add_bullet("API Endpoint POST /api/chat: Nhận message (tin nhắn hiện tại) và history (mảng lịch sử chat) từ Frontend.")
add_bullet("Chuyển đổi history sang định dạng Gemini (role: user/model) rồi gọi model.startChat() để duy trì ngữ cảnh hội thoại.")
add_bullet("Xử lý lỗi 429 (Too Many Requests): Khi vượt giới hạn 5 RPM hoặc 20 RPD, trả về thông báo thân thiện tiếng Việt thay vì lỗi kỹ thuật.")

add_heading_custom("3.2. Frontend (React + Vite)", size=13, space_before=6)

add_body("Cấu trúc thư mục:", bold=True, indent=0.5, space_after=2)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Cm(1)
p.paragraph_format.space_after = Pt(4)
add_run(p, "frontend/\n├── .dockerignore\n├── Dockerfile          (Multi-stage build)\n├── nginx.conf          (Reverse proxy config)\n├── vite.config.js      (Dev proxy /api → localhost:5000)\n├── index.html\n├── src/\n│   ├── App.jsx         (Component chính)\n│   ├── App.css         (Styling dark theme)\n│   └── main.jsx\n└── package.json", size=11)

add_body("Các tính năng chính trong App.jsx:", bold=True, indent=0.5, space_after=2)
add_bullet("State Management: Sử dụng useState để quản lý mảng messages (lưu cả tin nhắn user và bot), biến input và trạng thái loading.")
add_bullet("Gửi tin nhắn: Gọi API POST /api/chat với message hiện tại và toàn bộ history để AI hiểu ngữ cảnh.")
add_bullet("Giao diện chat: Phân biệt rõ tin nhắn \"Bạn\" (bên phải, gradient tím) và \"Bot\" (bên trái, xám đậm) với avatar riêng.")
add_bullet("Hiệu ứng Loading: Hiển thị animation \"Đang suy nghĩ...\" với 3 dấu chấm nhấp nháy khi chờ API phản hồi.")
add_bullet("Validation: Không cho phép gửi tin nhắn trống, disable nút Gửi và input khi đang loading.")
add_bullet("Auto-scroll: Tự động cuộn xuống tin nhắn mới nhất bằng useRef + scrollIntoView.")
add_bullet("Responsive Design: Giao diện dark theme hiện đại, tương thích tốt trên cả desktop và mobile.")

add_heading_custom("3.3. Docker & Docker Compose", size=13, space_before=6)

add_body("Cấu hình Docker:", bold=True, indent=0.5, space_after=2)

# Docker config table
table = doc.add_table(rows=4, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Table Grid'

headers = ["Thành phần", "Base Image", "Port", "Vai trò"]
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cell.paragraphs[0].add_run(h)
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = 'Times New Roman'

data = [
    ["Backend Dockerfile", "node:20-alpine", "5000", "Express API Server + Gemini AI SDK"],
    ["Frontend Dockerfile", "node:20-alpine → nginx:alpine", "80", "Multi-stage: Build React → Serve bằng Nginx"],
    ["docker-compose.yml", "—", "5000, 80", "Điều phối 2 services, truyền biến môi trường"],
]
for r, row_data in enumerate(data):
    for c, val in enumerate(row_data):
        cell = table.rows[r + 1].cells[c]
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(val)
        run.font.size = Pt(11)
        run.font.name = 'Times New Roman'

doc.add_paragraph()
add_body("Quy trình triển khai:", bold=True, indent=0.5, space_after=2)
add_bullet("Biến môi trường GEMINI_API_KEY được khai báo trong file .env ở thư mục gốc, Docker Compose tự động đọc và truyền vào container backend – không hardcode trong Dockerfile hay source code.")
add_bullet("Frontend container sử dụng nginx.conf tùy chỉnh: serve file tĩnh React tại / và reverse proxy mọi request /api/* sang http://backend:5000 thông qua Docker internal network.")
add_bullet("Service frontend được cấu hình depends_on: backend để đảm bảo backend khởi động trước.")
add_bullet("Lệnh duy nhất để build và chạy toàn bộ: docker compose up --build")

# ─── 4. Chức năng chatbot ──────────────────────────────────────────
doc.add_paragraph()
add_heading_custom("4. Các chức năng chính của Chatbot", size=13)

# Features table
table2 = doc.add_table(rows=7, cols=3)
table2.alignment = WD_TABLE_ALIGNMENT.CENTER
table2.style = 'Table Grid'

headers2 = ["STT", "Chức năng", "Mô tả chi tiết"]
for i, h in enumerate(headers2):
    cell = table2.rows[0].cells[i]
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cell.paragraphs[0].add_run(h)
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = 'Times New Roman'

features = [
    ["1", "Chat với AI", "Gửi tin nhắn và nhận phản hồi thông minh từ Gemini 2.5 Flash, trả lời ngắn gọn, lịch sự, đúng trọng tâm."],
    ["2", "Tự động dịch", "Khi nhận tin nhắn không phải tiếng Việt, AI tự động dịch sang tiếng Việt và đánh dấu [Dịch]."],
    ["3", "Tóm tắt hội thoại", "Gõ từ khóa \"tóm tắt\" để AI tổng hợp toàn bộ lịch sử trò chuyện thành các gạch đầu dòng ngắn gọn."],
    ["4", "Lịch sử ngữ cảnh", "Toàn bộ lịch sử chat được gửi kèm mỗi request, giúp AI hiểu ngữ cảnh xuyên suốt cuộc hội thoại."],
    ["5", "Xử lý giới hạn API", "Bắt lỗi 429 (Rate Limit) và hiển thị thông báo thân thiện: \"Bạn đã gửi quá nhiều yêu cầu...\""],
    ["6", "Giao diện hiện đại", "Dark theme, responsive, phân biệt tin nhắn user/bot bằng màu sắc và vị trí, animation loading."],
]
for r, row_data in enumerate(features):
    for c, val in enumerate(row_data):
        cell = table2.rows[r + 1].cells[c]
        if c == 0:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(val)
        run.font.size = Pt(11)
        run.font.name = 'Times New Roman'

# Set column widths
for row in table2.rows:
    row.cells[0].width = Cm(1.2)
    row.cells[1].width = Cm(3.5)
    row.cells[2].width = Cm(12)

# ─── 5. Khảo sát AI ───────────────────────────────────────────────
doc.add_paragraph()
add_heading_custom("5. Khảo sát chi tiết công nghệ AI (Dựa trên tài liệu tiếng Anh)", size=13)

add_body("Công nghệ: Gemini 2.5 Flash – Large Language Model của Google DeepMind", bold=True, indent=0.5)
add_body("Dựa trên tài liệu chính thức từ Google AI:", italic=True, indent=0.5)

add_bullet("Gemini được thiết kế để hiểu đồng thời Hình ảnh, Âm thanh và Văn bản. Trong dự án này, tính năng xử lý văn bản đa ngôn ngữ được tận dụng để tự động nhận diện và dịch tin nhắn.", bold_prefix="Mô hình Multimodal: ")
add_bullet("Tham số systemInstruction khi khởi tạo model cho phép định nghĩa \"nhân cách\" và quy tắc hoạt động của AI ngay từ đầu, đảm bảo mọi phản hồi đều tuân theo kịch bản đã thiết lập.", bold_prefix="System Instruction: ")
add_bullet("API hỗ trợ truyền mảng history với role (user/model), giúp AI duy trì ngữ cảnh hội thoại qua nhiều lượt trao đổi – nền tảng cho tính năng \"tóm tắt\".", bold_prefix="Multi-turn Conversation: ")
add_bullet("Tài liệu chính thức nhấn mạnh việc sử dụng Safety Settings để lọc nội dung không phù hợp – một cân nhắc quan trọng khi triển khai chatbot trong môi trường thực tế.", bold_prefix="Safety Settings: ")
add_bullet("\"Get started with the Gemini API\" (ai.google.dev).", bold_prefix="Tài liệu tham khảo: ")

# ─── 6. Kết luận ──────────────────────────────────────────────────
doc.add_paragraph()
add_heading_custom("6. Kết luận", size=13)

add_body("Dự án Smart AI Chatbot đã hoàn thành đầy đủ các mục tiêu đề ra:", indent=0.5)
add_bullet("Xây dựng thành công ứng dụng Full-Stack với kiến trúc Backend (Node.js/Express) – Frontend (React/Vite) rõ ràng.")
add_bullet("Tích hợp thành công Gemini 2.5 Flash API với 3 chức năng cốt lõi: tự động dịch, phản hồi thông minh, và tóm tắt hội thoại.")
add_bullet("Xử lý hoàn chỉnh các edge case: tin nhắn trống, lỗi 429 Rate Limit, lỗi server.")
add_bullet("Container hóa toàn bộ dự án bằng Docker Compose, sẵn sàng triển khai production với một lệnh duy nhất.")
add_bullet("Giao diện responsive, dark theme hiện đại, trải nghiệm người dùng mượt mà với loading animation và auto-scroll.")

# ─── Save ──────────────────────────────────────────────────────────
output_path = r"D:\smart-ai\BaoCao2_ThucTapCoSo_Nhom7_LyHaiQuan.docx"
doc.save(output_path)
print(f"✅ Đã tạo báo cáo tại: {output_path}")
